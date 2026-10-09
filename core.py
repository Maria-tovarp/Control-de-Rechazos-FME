from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
from io import BytesIO
import re
import unicodedata
from urllib.parse import quote
from typing import Dict, List

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


def norm(value):
    if value is None:
        return ""
    return str(value).strip()


def normalize_header(value):
    return re.sub(r"\s+", " ", norm(value)).casefold()


def normalize_site(value):
    return re.sub(r"\s+", "", norm(value)).upper()




def normalize_digits(value):
    return re.sub(r"\D+", "", norm(value))


def normalize_person_name(value):
    text = unicodedata.normalize("NFKD", norm(value))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^A-Za-z0-9]+", " ", text).upper()
    return re.sub(r"\s+", " ", text).strip()


def extract_fme_document(fme_value: str) -> str:
    text = norm(fme_value)
    m = re.search(r"\((\d{6,12})\)\s*$", text)
    if m:
        return m.group(1)
    # Si el campo viene solo como cédula/ID, usarlo directamente.
    digits = normalize_digits(text)
    if text and digits == text.replace(" ", "") and 6 <= len(digits) <= 12:
        return digits
    return ""


def extract_fme_name(fme_value: str) -> str:
    text = re.sub(r"\(\d{6,12}\)\s*$", "", norm(fme_value)).strip()
    return normalize_person_name(text)


def normalize_mobile(value) -> str:
    digits = normalize_digits(value)
    if not digits:
        return ""
    if digits.startswith("57") and len(digits) == 12:
        return digits
    if len(digits) == 10 and digits.startswith("3"):
        return "57" + digits
    return digits


def load_fme_directory(cuadrillas_bytes: bytes) -> Dict[str, dict]:
    """Construye directorio FME desde CUADRILLAS ACT. NUEVO PROY. 2026.xlsx.

    Busca por CEDULA como llave principal y por nombre normalizado como respaldo.
    Se leen todas las hojas con columnas NOMBRE/NOMBRES, CEDULA y CELULAR.
    """
    wb = load_workbook(BytesIO(cuadrillas_bytes), data_only=True)
    by_doc, by_name = {}, {}
    for ws in wb.worksheets:
        headers = [normalize_header(c.value) for c in ws[1]]
        name_idx = next((i for i,h in enumerate(headers) if h in {"nombre", "nombres"}), None)
        doc_idx = next((i for i,h in enumerate(headers) if h in {"cedula", "cédula"}), None)
        phone_idx = next((i for i,h in enumerate(headers) if h in {"celular", "telefono", "teléfono"}), None)
        if name_idx is None or doc_idx is None or phone_idx is None:
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            name = norm(row[name_idx] if name_idx < len(row) else None)
            doc = normalize_digits(row[doc_idx] if doc_idx < len(row) else None)
            raw_phone = norm(row[phone_idx] if phone_idx < len(row) else None)
            phone = normalize_mobile(raw_phone)
            if not name and not doc:
                continue
            rec = {
                "Nombre FME data": name,
                "Cedula FME": doc,
                "Celular FME": phone,
                "Celular original": raw_phone,
                "Origen celular": ws.title,
            }
            if doc and doc not in by_doc:
                by_doc[doc] = rec
            nkey = normalize_person_name(name)
            if nkey and nkey not in by_name:
                by_name[nkey] = rec
    return {"__by_doc__": by_doc, "__by_name__": by_name}


def find_fme_contact(fme_value: str, directory: Dict[str, dict]) -> dict:
    doc = extract_fme_document(fme_value)
    if doc and doc in directory.get("__by_doc__", {}):
        rec = dict(directory["__by_doc__"][doc])
        rec["Origen match celular"] = "CEDULA"
        return rec
    name = extract_fme_name(fme_value)
    if name and name in directory.get("__by_name__", {}):
        rec = dict(directory["__by_name__"][name])
        rec["Origen match celular"] = "NOMBRE"
        return rec
    return {
        "Nombre FME data": "", "Cedula FME": doc, "Celular FME": "",
        "Celular original": "", "Origen celular": "No encontrado", "Origen match celular": "No encontrado"
    }

def whatsapp_url(phone: str, message: str) -> str:
    phone = normalize_mobile(phone)
    if not phone:
        return ""
    return f"https://wa.me/{phone}?text={quote(message)}"


def to_datetime(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, (int, float)):
        try:
            return from_excel(value)
        except Exception:
            return None
    text = norm(value)
    for fmt in (
        "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y",
        "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    return None


def worksheet_records(ws) -> List[dict]:
    headers = [norm(c.value) for c in ws[1]]
    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec, nonempty = {}, False
        for i, header in enumerate(headers):
            if not header:
                continue
            value = row[i] if i < len(row) else None
            rec[header] = value
            if value not in (None, ""):
                nonempty = True
        if nonempty:
            records.append(rec)
    return records


def get_field(record: dict, *names):
    index = {normalize_header(k): k for k in record.keys()}
    for name in names:
        key = index.get(normalize_header(name))
        if key is not None:
            return record.get(key)
    return None


def extract_existing_assignment_ids(base_wb) -> set:
    ids = set()
    for ws in base_wb.worksheets:
        if ws.max_row < 2:
            continue
        headers = [normalize_header(c.value) for c in ws[1]]
        if normalize_header("Assignment ID") not in headers:
            continue
        col = headers.index(normalize_header("Assignment ID")) + 1
        for r in range(2, ws.max_row + 1):
            val = norm(ws.cell(r, col).value)
            if val:
                ids.add(val)
    return ids


def learn_coordinator_map(base_wb) -> Dict[str, str]:
    votes = defaultdict(Counter)
    for ws in base_wb.worksheets:
        if ws.max_row < 2:
            continue
        headers = [norm(c.value) for c in ws[1]]
        hnorm = [normalize_header(h) for h in headers]
        coord_idx = fme_idx = None
        for candidate in ("Coordinador", "Coordinador reporte"):
            if normalize_header(candidate) in hnorm:
                coord_idx = hnorm.index(normalize_header(candidate)) + 1
                break
        for candidate in ("FME ID", "FME", "FME "):
            if normalize_header(candidate) in hnorm:
                fme_idx = hnorm.index(normalize_header(candidate)) + 1
                break
        if not coord_idx or not fme_idx:
            continue
        for r in range(2, ws.max_row + 1):
            coord = norm(ws.cell(r, coord_idx).value)
            fme = norm(ws.cell(r, fme_idx).value)
            if not fme or fme in {"()", "#N/A", "Sin FME"}:
                continue
            if not coord or coord in {"#N/A", "Sin coordinador", "Por validar"}:
                continue
            votes[fme][coord] += 1
    return {fme: counts.most_common(1)[0][0] for fme, counts in votes.items()}


def load_site_directory(coordinator_bytes: bytes) -> Dict[str, dict]:
    """Replica exactamente el BUSCARX usado en Excel.

    Prioridad:
    1) Hoja2!D:D (CODIGO)       -> Hoja2!B:B (Coordinador)
    2) Hoja2!E:E (CodigoNetco)  -> Hoja2!B:B (Coordinador)
    3) Hoja1!A:A (SITIO)        -> Hoja1!B:B (COORDINADOR)

    No usa Otros Códigos ni infiere el coordinador desde el histórico.
    """
    wb = load_workbook(BytesIO(coordinator_bytes), data_only=True)

    hoja2_d: Dict[str, dict] = {}
    hoja2_e: Dict[str, dict] = {}
    hoja1_a: Dict[str, dict] = {}

    if "Hoja2" in wb.sheetnames:
        ws = wb["Hoja2"]
        # Se usan posiciones de columna exactas para reproducir la fórmula del usuario:
        # B=Coordinador, D=CODIGO, E=CodigoNetco.
        for row in ws.iter_rows(min_row=2, values_only=True):
            coord = norm(row[1] if len(row) > 1 else None)
            codigo = norm(row[3] if len(row) > 3 else None)
            codigo_netco = norm(row[4] if len(row) > 4 else None)
            if not coord:
                continue
            data = {
                "Coordinador": coord,
                "Zona": norm(row[0] if len(row) > 0 else None),
                "Codigo": codigo,
                "CodigoNetco": codigo_netco,
                "Origen coordinador": "Hoja2 - CODIGO (columna D)",
            }
            kd = normalize_site(codigo)
            if kd and kd not in {"-", "#N/A", "NOAPLICA"} and kd not in hoja2_d:
                hoja2_d[kd] = data

            ke = normalize_site(codigo_netco)
            if ke and ke not in {"-", "#N/A", "NOAPLICA"} and ke not in hoja2_e:
                data_e = dict(data)
                data_e["Origen coordinador"] = "Hoja2 - CodigoNetco (columna E)"
                hoja2_e[ke] = data_e

    if "Hoja1" in wb.sheetnames:
        ws = wb["Hoja1"]
        # A=SITIO, B=COORDINADOR.
        for row in ws.iter_rows(min_row=2, values_only=True):
            site = norm(row[0] if len(row) > 0 else None)
            coord = norm(row[1] if len(row) > 1 else None)
            key = normalize_site(site)
            if key and coord and key not in hoja1_a:
                hoja1_a[key] = {
                    "Coordinador": coord,
                    "Origen coordinador": "Hoja1 - SITIO (columna A)",
                }

    # Guardamos los tres índices por separado. find_site_coordinator() aplica la prioridad.
    return {"__hoja2_d__": hoja2_d, "__hoja2_e__": hoja2_e, "__hoja1_a__": hoja1_a}


def find_site_coordinator(site_id: str, site_directory: Dict[str, dict]) -> dict:
    key = normalize_site(site_id)
    for bucket in ("__hoja2_d__", "__hoja2_e__", "__hoja1_a__"):
        data = site_directory.get(bucket, {}).get(key)
        if data:
            return data
    return {"Coordinador": "Por validar", "Origen coordinador": "No encontrado en DATA COORDINADORES"}


def find_report_sheet(report_wb):
    preferred = ["WO Result query Custom", "WO Result query"]
    for name in preferred:
        if name in report_wb.sheetnames:
            return report_wb[name]
    for ws in report_wb.worksheets:
        headers = {normalize_header(c.value) for c in ws[1]}
        required = {normalize_header(x) for x in ["Assignment ID", "Audit time", "FME ID", "Audit remark", "Audit status"]}
        if required.issubset(headers):
            return ws
    raise ValueError("No encontré una hoja con las columnas esperadas del reporte de auditoría.")


CATEGORY_RULES = {
    "Evidencia fotográfica": [
        "foto", "fotograf", "imagen", "evidencia", "no se observa", "no se visualiza", "debe verse", "verse completo", "registro fotográfico"
    ],
    "Intervención física": [
        "retirar", "quitar", "abrir", "tapa", "breaker", "rack", "rectificador", "modulo", "módulo", "puerta", "soporte", "instalar", "ajustar"
    ],
    "Limpieza / organización": [
        "limpieza", "limpiar", "basura", "sobrante", "desmalez", "organizado", "organización", "sellamiento"
    ],
    "Edición de formulario / texto": [
        "eliminar formulario", "borrar formulario", "quitar formulario", "modificar formulario",
        "agregar texto", "añadir texto", "adicionar texto", "agregar información", "añadir información",
        "diligenciar", "corregir texto", "modificar texto", "editar formulario", "editar campo"
    ],
    "Documentación / certificación": [
        "certific", "document", "soporte documental", "acta", "formato"
    ],
    "Información / códigos": [
        "codigo", "código", "serial", "marquilla", "referencia", "id del rack", "diligenciar", "información", "micode", "azimut"
    ],
    "Hallazgo / condición técnica": [
        "hallazgo", "alarma", "aterriz", "tierra", "voltaje", "corriente", "bater", "dañado", "falla"
    ],
}


def classify_audit_remark(remark: str) -> List[str]:
    text = norm(remark).casefold()
    categories = []
    for category, words in CATEGORY_RULES.items():
        if any(w.casefold() in text for w in words):
            categories.append(category)
    return categories or ["Revisión manual"]


def determine_closed_management(categories: List[str]) -> str:
    cats = set(categories)
    has_images = "Evidencia fotográfica" in cats
    has_noneditable = bool(cats & {
        "Edición de formulario / texto", "Intervención física", "Limpieza / organización",
        "Documentación / certificación", "Información / códigos", "Hallazgo / condición técnica"
    })
    if has_images and not has_noneditable:
        return "Enviar evidencias por correo"
    if has_images and has_noneditable:
        return "Correo mixto: evidencias + solicitud de exclusión"
    return "Solicitud de exclusión"


def build_purpose(categories: List[str]) -> str:
    cats = set(categories)
    has_images = "Evidencia fotográfica" in cats
    noneditable = []

    if "Edición de formulario / texto" in cats:
        noneditable.append("eliminar o modificar formularios, agregar texto o actualizar información dentro de la actividad")
    if "Intervención física" in cats:
        noneditable.append("realizar verificaciones o ajustes físicos sobre los equipos o elementos del sitio")
    if "Limpieza / organización" in cats:
        noneditable.append("ejecutar actividades físicas de limpieza, organización o adecuación")
    if "Documentación / certificación" in cats:
        noneditable.append("complementar documentación o soportes dentro de la actividad")
    if "Información / códigos" in cats:
        noneditable.append("actualizar o recapturar información técnica, códigos, seriales o marquillas dentro de la actividad")
    if "Hallazgo / condición técnica" in cats:
        noneditable.append("realizar verificaciones o correcciones técnicas en sitio")

    if has_images and not noneditable:
        return (
            "El propósito de este mensaje es atender el requerimiento de auditoría relacionado con evidencias fotográficas. "
            "Dado que la actividad se encuentra en estado Closed, las imágenes solicitadas no pueden incorporarse nuevamente dentro de la actividad; "
            "sin embargo, sí pueden ser remitidas por correo electrónico para su validación. "
            "Por lo anterior, se envían por este medio las evidencias solicitadas para su revisión."
        )

    if noneditable:
        actions = noneditable[0] if len(noneditable) == 1 else ", ".join(noneditable[:-1]) + " y " + noneditable[-1]
    else:
        actions = "realizar las correcciones solicitadas dentro de la actividad"

    if has_images:
        return (
            "El propósito de este mensaje es atender los hallazgos identificados durante la auditoría. "
            "Las evidencias fotográficas solicitadas pueden ser remitidas por correo electrónico y se adjuntan por este medio para su validación. "
            f"No obstante, el rechazo también requiere {actions}. "
            "Dado que la actividad se encuentra en estado Closed, no es posible efectuar dichas modificaciones o correcciones dentro del sistema. "
            "Por lo anterior, agradecemos gestionar la exclusión correspondiente para los aspectos que no pueden ser corregidos en la actividad cerrada, "
            "con el fin de evitar un impacto negativo en la medición de los KPI."
        )

    return (
        "El propósito de este mensaje es solicitar formalmente la exclusión de la actividad en referencia. "
        f"Los hallazgos identificados durante la auditoría requieren {actions}. "
        "Dado que la actividad se encuentra en estado Closed, no es posible realizar estas modificaciones o correcciones dentro del sistema. "
        "Por lo anterior, agradecemos gestionar la exclusión correspondiente con el fin de evitar un impacto negativo en la medición de los KPI."
    )


def detect_rejections(report_wb, year: int, month: int, existing_ids: set, coordinator_map: Dict[str, str], site_directory: Dict[str, dict], fme_directory: Dict[str, dict], feedbacks: Dict[str, dict] | None = None):
    feedbacks = feedbacks or {}
    ws = find_report_sheet(report_wb)
    records = worksheet_records(ws)
    out, seen = [], set()
    for rec in records:
        status = norm(get_field(rec, "Audit status"))
        if status.casefold() != "reject":
            continue
        audit_dt = to_datetime(get_field(rec, "Audit time"))
        if not audit_dt or audit_dt.year != year or audit_dt.month != month:
            continue
        assignment = norm(get_field(rec, "Assignment ID"))
        if not assignment or assignment in seen:
            continue
        seen.add(assignment)
        fme = norm(get_field(rec, "FME ID")) or "Sin FME"
        site = norm(get_field(rec, "Site ID"))
        task_id = norm(get_field(rec, "Task ID"))
        site_data = find_site_coordinator(site, site_directory)
        fme_data = find_fme_contact(fme, fme_directory)
        coord = site_data.get("Coordinador", "Por validar")
        remark = norm(get_field(rec, "Audit remark"))
        categories = classify_audit_remark(remark)
        task_status = norm(get_field(rec, "Task status"))
        fb = feedbacks.get(assignment) or feedbacks.get(task_id) or {}
        out.append({
            "Coordinador": coord,
            "Assignment ID": assignment,
            "Create time": get_field(rec, "Create time"),
            "Assign time": get_field(rec, "Assign time"),
            "Assign operator": get_field(rec, "Assign operator"),
            "Audit time": audit_dt,
            "Task ID": task_id,
            "Task status": task_status,
            "Site ID": site,
            "Region": norm(get_field(rec, "Region")),
            "FM Office": norm(get_field(rec, "FM Office")),
            "FME ID": fme,
            "FME Supplier": norm(get_field(rec, "FME Supplier")),
            "Audit remark": remark,
            "Audit status": status,
            "Feedback Remark": norm(get_field(rec, "Feedback Remark")),
            "Feedback Qty": get_field(rec, "Feedback Qty"),
            "Audit type": norm(get_field(rec, "Audit type")),
            "Customer ticket": norm(get_field(rec, "Customer ticket")),
            "Task title": norm(get_field(rec, "Task title")),
            "Assign User": norm(get_field(rec, "Assign User")),
            "Estado contacto": "Pendiente",
            "Es nuevo": "Sí" if assignment not in existing_ids else "No",
            "Es Closed": task_status.casefold() == "closed",
            "Tipo rechazo": " + ".join(categories),
            "Categorias": categories,
            "Proposito": build_purpose(categories),
            "Gestion Closed": determine_closed_management(categories),
            "Departamento": site_data.get("Departamento", ""),
            "Municipio": site_data.get("Municipio", ""),
            "Nombre sitio": site_data.get("Nombre sitio", ""),
            "Origen coordinador": site_data.get("Origen coordinador", ""),
            "Celular FME": fme_data.get("Celular FME", ""),
            "Celular original": fme_data.get("Celular original", ""),
            "Origen celular": fme_data.get("Origen celular", ""),
            "Origen match celular": fme_data.get("Origen match celular", ""),
            "Nombre FME data": fme_data.get("Nombre FME data", ""),
            "Fecha de feedback": norm(fb.get("fecha")),
            "Estado OWS": norm(fb.get("estado")),
            "Feedback OWS": norm(fb.get("feedback")),
            "Tiene feedback OWS": "Sí" if norm(fb.get("fecha")) and norm(fb.get("estado")) and norm(fb.get("feedback")) else "No",
        })
    return out

def fme_first_name(fme: str) -> str:
    text = re.sub(r"\([^)]*\)\s*$", "", norm(fme)).strip()
    return text.split()[0].title() if text else ""


def build_individual_message(row: dict) -> str:
    name = fme_first_name(row.get("FME ID", ""))
    greet = f"Buen día, {name}." if name else "Buen día."
    return (
        f"{greet} Se presenta un rechazo de auditoría en el sitio {row.get('Site ID','')} "
        f"(Task ID {row.get('Task ID','')}).\n\n"
        f"Audit Remark:\n{row.get('Audit remark','')}\n\n"
        "Por favor realizar las correcciones indicadas y confirmar cuando quede gestionado. Gracias."
    )


def group_messages(rows: List[dict]) -> List[dict]:
    # No se pide corrección al FME si la actividad ya está Closed.
    rows = [r for r in rows if not r.get("Es Closed") and r.get("Es nuevo") == "Sí"]
    grouped = defaultdict(list)
    for row in rows:
        grouped[row.get("FME ID", "Sin FME")].append(row)
    result = []
    for fme, items in sorted(grouped.items(), key=lambda x: x[0].casefold()):
        source_name = next((r.get("Nombre FME data", "") for r in items if r.get("Nombre FME data")), "")
        name = source_name.split()[0].title() if source_name else fme_first_name(fme)
        greet = f"Buen día, {name}." if name else "Buen día."
        chunks = []
        for idx, r in enumerate(items, 1):
            chunks.append(
                f"{idx}. Sitio: {r.get('Site ID','')}\n"
                f"Task ID: {r.get('Task ID','')}\n"
                f"Audit Remark: {r.get('Audit remark','')}"
            )
        message = (
            f"{greet} Te comparto los rechazos de auditoría pendientes de corrección:\n\n"
            + "\n\n".join(chunks)
            + "\n\nPor favor realizar las correcciones indicadas y confirmar cuando queden gestionadas. Gracias."
        )
        phone = next((r.get("Celular FME", "") for r in items if r.get("Celular FME")), "")
        contact_keys = [norm(r.get("Assignment ID") or r.get("Task ID")) for r in items]
        result.append({
            "Coordinador": items[0].get("Coordinador", "Por validar"),
            "FME": fme,
            "Cantidad rechazos": len(items),
            "Sitios": ", ".join(sorted({r.get('Site ID','') for r in items if r.get('Site ID')})),
            "Mensaje": message,
            "Estado contacto": "Pendiente",
            "Claves contacto": [key for key in contact_keys if key],
            "Celular FME": phone,
            "Celular visible": phone[2:] if phone.startswith("57") and len(phone) == 12 else phone,
            "WhatsApp URL": whatsapp_url(phone, message),
            "Origen celular": next((r.get("Origen celular", "") for r in items if r.get("Celular FME")), "No encontrado"),
        })
    return result

def format_dt_es(dt: datetime) -> str:
    if not dt:
        return ""
    hour = dt.hour
    suffix = "a. m." if hour < 12 else "p. m."
    hour12 = hour % 12 or 12
    return f"{dt:%d/%m/%Y} {hour12}:{dt:%M:%S} {suffix}"


def build_closed_email(row: dict, close_date: str) -> dict:
    close_dt = to_datetime(close_date)
    close_text = close_dt.strftime("%Y-%m-%d %H:%M:%S") if close_dt else norm(close_date)
    audit_text = format_dt_es(row.get("Audit time"))
    task_id = row.get("Task ID") or row.get("Assignment ID")
    management = row.get("Gestion Closed", determine_closed_management(row.get("Categorias", [])))
    if management == "Enviar evidencias por correo":
        subject = f"ENVÍO DE EVIDENCIAS PARA {task_id} EN ESTADO CLOSED"
    elif management.startswith("Correo mixto"):
        subject = f"EVIDENCIAS Y SOLICITUD DE EXCLUSIÓN PARA {task_id} EN ESTADO CLOSED"
    else:
        subject = f"EXCLUSIÓN PARA {task_id} EN ESTADO CLOSED"
    body = (
        "Buenos días, cordial saludo.\n\n"
        "De acuerdo a lo mencionado con anterioridad con respecto al cierre de actividades, "
        f"esta fue cerrada automáticamente por el sistema el día {close_text}. Posteriormente, "
        f"fue auditada por parte de GNOC, {audit_text}. Quedando en un estado rechazada.\n\n"
        f"{task_id}\n"
        "closed\n"
        f"{row.get('Site ID','')}\n"
        f"{row.get('Audit time').strftime('%d/%m/%Y') if row.get('Audit time') else ''} Reject GNOC\n"
        f"{row.get('Audit remark','')}\n\n"
        f"{row.get('Proposito','')}"
    )
    return {"Asunto": subject, "Cuerpo": body}


def _safe_sheet_name(name: str) -> str:
    return re.sub(r"[\\/*?:\[\]]", "-", name)[:31]


def _style_new_sheet(ws):
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def update_workbook(base_bytes: bytes, rows: List[dict], year: int, month: int) -> bytes:
    wb = load_workbook(BytesIO(base_bytes))
    month_names = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]
    label = month_names[month - 1]
    detail_name = _safe_sheet_name(f"RECHAZOS {label}")
    msg_name = _safe_sheet_name(f"MENSAJES FME {label}")
    closed_name = _safe_sheet_name(f"CLOSED {label}")

    new_rows = [r for r in rows if r.get("Es nuevo") == "Sí"]
    if detail_name in wb.sheetnames:
        del wb[detail_name]
    ws = wb.create_sheet(detail_name)
    headers = [
        "Coordinador", "Origen coordinador", "Assignment ID", "Create time", "Assign time", "Assign operator", "Audit time",
        "Task ID", "Task status", "Site ID", "Region", "FM Office", "FME ID", "FME Supplier",
        "Audit remark", "Audit status", "Feedback Remark", "Feedback Qty", "Audit type", "Customer ticket",
        "Task title", "Assign User", "Estado contacto", "Tipo rechazo", "Departamento", "Municipio", "Nombre sitio",
        "Celular FME", "Fecha de feedback", "Estado OWS", "Feedback OWS", "Tiene feedback OWS"
    ]
    ws.append(headers)
    for r in new_rows:
        ws.append([r.get(h, "") for h in headers])
    _style_new_sheet(ws)
    for c in ws["F"][1:]:
        c.number_format = "dd/mm/yyyy hh:mm"
    for col in ("A","B","F","G","H","I","L","O","V","W","X","Y","Z"):
        ws.column_dimensions[col].width = 20
    ws.column_dimensions["N"].width = 80

    if msg_name in wb.sheetnames:
        del wb[msg_name]
    mw = wb.create_sheet(msg_name)
    mh = ["Coordinador", "FME", "Celular FME", "Cantidad rechazos", "Sitios", "Mensaje", "Estado contacto"]
    mw.append(mh)
    for item in group_messages(rows):
        mw.append([item.get(h, "") for h in mh])
    _style_new_sheet(mw)
    mw.column_dimensions["E"].width = 100

    if closed_name in wb.sheetnames:
        del wb[closed_name]
    cw = wb.create_sheet(closed_name)
    ch = ["Coordinador", "Origen coordinador", "Task ID", "Assignment ID", "Site ID", "FME ID", "Audit time", "Audit remark", "Tipo rechazo", "Gestion Closed", "Fecha cierre OWS", "Estado gestión"]
    cw.append(ch)
    for r in rows:
        if r.get("Es Closed"):
            cw.append([r.get(h, "") for h in ch[:-2]] + ["", "Falta fecha OWS"])
    _style_new_sheet(cw)
    cw.column_dimensions["G"].width = 90
    cw.column_dimensions["H"].width = 45
    cw.column_dimensions["I"].width = 22
    cw.column_dimensions["J"].width = 22

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def process_files(base_bytes: bytes, report_bytes: bytes, coordinator_bytes: bytes, cuadrillas_bytes: bytes, year: int, month: int, feedbacks: Dict[str, dict] | None = None):
    base_wb = load_workbook(BytesIO(base_bytes), data_only=False)
    report_wb = load_workbook(BytesIO(report_bytes), data_only=True)
    existing = extract_existing_assignment_ids(base_wb)
    coordinator_map = {}  # El coordinador sale únicamente de DATA 2026 - COORDINADORES.xlsx
    site_directory = load_site_directory(coordinator_bytes)
    fme_directory = load_fme_directory(cuadrillas_bytes)
    rows = detect_rejections(report_wb, year, month, existing, coordinator_map, site_directory, fme_directory, feedbacks=feedbacks)
    messages = group_messages(rows)
    output = update_workbook(base_bytes, rows, year, month)
    return rows, messages, output, site_directory, fme_directory
